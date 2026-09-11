import datetime
import json
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')


def clean(value):
    if value is None or value == '-':
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def date_value(value):
    value = clean(value)
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime('%Y-%m-%d')
    if value:
        return str(value)[:10]
    return None


def number(value):
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace('.', '').replace(',', '.'))
    except ValueError:
        return None


def headers_in(sheet):
    for row in range(1, sheet.max_row + 1):
        values = [clean(sheet.cell(row, column).value) for column in range(1, sheet.max_column + 1)]
        if 'Data Entrada' in values and 'Ativo' in values and 'Status' in values:
            return row, {value: index + 1 for index, value in enumerate(values) if value}
    return None, None


def field(sheet, row, mapping, name):
    return clean(sheet.cell(row, mapping[name]).value) if name in mapping else None


def main(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    rows = []
    peels = []
    peel_sheet = workbook['PEEL_OFF_LOG'] if 'PEEL_OFF_LOG' in workbook.sheetnames else None
    if peel_sheet:
        for row in range(6, peel_sheet.max_row + 1):
            origin = clean(peel_sheet.cell(row, 2).value)
            entry = date_value(peel_sheet.cell(row, 3).value)
            ticker = clean(peel_sheet.cell(row, 4).value)
            price = number(peel_sheet.cell(row, 6).value)
            quantity = number(peel_sheet.cell(row, 8).value)
            exit_price = number(peel_sheet.cell(row, 9).value)
            if origin and entry and ticker and price is not None and quantity and exit_price is not None:
                peels.append({'origin': origin, 'entryDate': entry, 'ticker': str(ticker).upper(), 'entryPrice': price, 'quantity': quantity, 'price': exit_price, 'note': clean(peel_sheet.cell(row, 12).value) or ''})
    for sheet in workbook.worksheets:
        header_row, columns = headers_in(sheet)
        if not columns:
            continue
        mode = 'paper' if 'Paper Trading' in sheet.title else 'real'
        for row in range(header_row + 1, sheet.max_row + 1):
            ticker = clean(field(sheet, row, columns, 'Ativo'))
            entry_date = date_value(field(sheet, row, columns, 'Data Entrada'))
            entry_price = number(field(sheet, row, columns, 'Entrada'))
            stop_price = number(field(sheet, row, columns, 'Stop Inicial'))
            quantity = number(field(sheet, row, columns, 'Qtd Inicial'))
            # The quarterly sheets retain formula and formatting rows below the actual
            # ledger. A row without an asset, price, or quantity is not a trade.
            if not any([ticker, entry_price, quantity]):
                continue
            record = {'sheet': sheet.title, 'row': row, 'mode': mode, 'ticker': str(ticker or '').upper(), 'entryDate': entry_date, 'exitDate': date_value(field(sheet, row, columns, 'Data Saída')), 'market': clean(field(sheet, row, columns, 'Mercado')), 'direction': clean(field(sheet, row, columns, 'Lado')), 'setup': clean(field(sheet, row, columns, 'Setup 1-2-3')), 'status': clean(field(sheet, row, columns, 'Status')), 'entryPrice': entry_price, 'stopPrice': stop_price, 'atr': number(field(sheet, row, columns, 'ATR (21)\nDiario')), 'exitPrice': number(field(sheet, row, columns, 'Atual/Saída')), 'quantity': quantity, 'riskPct': number(field(sheet, row, columns, 'Risco %')), 'financialResult': number(field(sheet, row, columns, 'P&L R$')), 'rMultiple': number(field(sheet, row, columns, 'R Múltiplo')), 'timeframe': clean(field(sheet, row, columns, 'Timeframe')), 'marketRegime': clean(field(sheet, row, columns, 'Regime Mercado')), 'riskPolicy': clean(field(sheet, row, columns, 'Política de Risco')), 'positionSizing': number(field(sheet, row, columns, 'Position Size Final')), 'limitedBy': clean(field(sheet, row, columns, 'Limitado Por')), 'executionScore': number(field(sheet, row, columns, 'Execution Score')), 'note': clean(field(sheet, row, columns, 'Observações / Padrões de Comportamento (Explique o motivo. Não deixe generico)')) or ''}
            record['peels'] = [item for item in peels if item['origin'] == sheet.title and item['entryDate'] == entry_date and item['ticker'] == record['ticker'] and abs(item['entryPrice'] - (entry_price or 0)) < 0.00001]
            rows.append(record)
    print(json.dumps({'trades': rows}, ensure_ascii=False, default=str))


if __name__ == '__main__':
    main(sys.argv[1])
